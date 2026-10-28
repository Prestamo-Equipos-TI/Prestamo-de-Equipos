from io import BytesIO

from django.contrib.auth.models import User
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from averias.models import ReporteAveria
from inventario.models import Equipo
from prestamos.models import SolicitudPrestamo
from usuarios.models import PerfilUsuario


class ReportesTest(TestCase):
    def setUp(self):
        self.admin = User.objects.create_user('adminqa', password='clave-segura')
        PerfilUsuario.objects.create(user=self.admin, rol='admin_ti', carnet='ADM-QA', nombre_completo='Admin QA')
        self.usuario = User.objects.create_user('usuarioqa', email='usuario@example.test', password='clave-segura')
        PerfilUsuario.objects.create(user=self.usuario, rol='usuario', carnet='USR-QA', nombre_completo='Usuario QA')
        self.equipo = Equipo.objects.create(codigo='EQ-QA-1', nombre='Laptop QA', categoria='laptop',
                                           marca='Marca QA', estado='disponible', ubicacion='Laboratorio')
        Equipo.objects.create(codigo='EQ-QA-2', nombre='Proyector QA', categoria='proyector',
                              marca='Marca QA', estado='mantenimiento', ubicacion='Bodega')
        self.solicitud = SolicitudPrestamo.objects.create(usuario=self.usuario, equipo=self.equipo,
                                                         fecha_prestamo='2026-10-27')
        ReporteAveria.objects.create(equipo=self.equipo, reportado_por=self.usuario,
                                     descripcion_falla='Pantalla con falla')
        self.client.force_login(self.admin)

    def test_inventory_export_respects_filters_and_keeps_text_safe(self):
        url = reverse('inventario:lista')
        respuesta = self.client.get(url, {'export': 'excel', 'estado': 'disponible', 'q': 'EQ-QA-1'})
        self.assertEqual(respuesta.status_code, 200)
        libro = load_workbook(BytesIO(respuesta.content))
        filas = list(libro.active.values)
        self.assertEqual(len(filas), 2)
        self.assertEqual(filas[1][0], 'EQ-QA-1')
        self.assertNotIn('EQ-QA-2', str(filas))
        self.equipo.nombre = '=1+1'
        self.equipo.save(update_fields=['nombre'])
        seguro = load_workbook(BytesIO(self.client.get(url, {'export': 'excel'}).content))
        celda = next(c for row in seguro.active.iter_rows(min_row=2) for c in row if c.value == '=1+1')
        self.assertEqual(celda.data_type, 's')
        pdf = self.client.get(url, {'export': 'pdf', 'estado': 'disponible', 'q': 'EQ-QA-1'})
        self.assertTrue(pdf.content.startswith(b'%PDF'))

    def test_every_report_downloads_and_pages_show_buttons(self):
        pages = [
            reverse('inventario:lista'),
            reverse('prestamos:solicitudes_admin'),
            reverse('historial:usuarios'),
            reverse('historial:usuario_detalle', args=[self.usuario.id]),
            reverse('historial:equipos'),
            reverse('historial:equipo_detalle', args=[self.equipo.id]),
        ]
        for url in pages:
            with self.subTest(url=url):
                page = self.client.get(url)
                self.assertContains(page, 'Reporte PDF')
                self.assertContains(page, 'Reporte Excel')
                pdf = self.client.get(url, {'export': 'pdf'})
                self.assertTrue(pdf.content.startswith(b'%PDF'))
                excel = self.client.get(url, {'export': 'excel'})
                self.assertEqual(excel.status_code, 200)
                self.assertIn('attachment;', excel['Content-Disposition'])
                self.assertGreater(len(load_workbook(BytesIO(excel.content)).sheetnames), 0)

    def test_report_content_matches_each_admin_screen(self):
        casos = [
            (reverse('prestamos:solicitudes_admin'), 'Solicitudes', 'Usuario QA'),
            (reverse('historial:usuarios'), 'Solicitudes', 'EQ-QA-1'),
            (reverse('historial:usuario_detalle', args=[self.usuario.id]), 'Solicitudes', 'EQ-QA-1'),
            (reverse('historial:equipos'), 'Averías', 'Pantalla con falla'),
            (reverse('historial:equipo_detalle', args=[self.equipo.id]), 'Averías', 'Pantalla con falla'),
        ]
        for url, hoja, valor in casos:
            with self.subTest(url=url):
                respuesta = self.client.get(url, {'export': 'excel'})
                libro = load_workbook(BytesIO(respuesta.content))
                self.assertIn(hoja, libro.sheetnames)
                self.assertIn(valor, [celda for fila in libro[hoja].values for celda in fila])

    def test_report_pages_are_admin_only(self):
        self.client.force_login(self.usuario)
        self.assertNotContains(self.client.get(reverse('prestamos:solicitar')), 'Reporte PDF')
        self.assertNotContains(self.client.get(reverse('prestamos:mis_solicitudes')), 'Reporte Excel')
        for url in (reverse('inventario:lista'), reverse('prestamos:solicitudes_admin'),
                    reverse('historial:usuarios'), reverse('historial:equipos')):
            with self.subTest(url=url):
                self.assertEqual(self.client.get(url).status_code, 403)
                self.assertEqual(self.client.get(url, {'export': 'excel'}).status_code, 403)

    def test_reports_require_login(self):
        self.client.logout()
        for url in (reverse('inventario:lista'), reverse('historial:equipos')):
            with self.subTest(url=url):
                self.assertEqual(self.client.get(url, {'export': 'pdf'}).status_code, 302)
