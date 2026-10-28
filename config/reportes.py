from datetime import date, datetime
from functools import wraps
from html import escape
from io import BytesIO

from django.http import HttpResponse, HttpResponseForbidden
from usuarios.models import PerfilUsuario
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle


def solo_admin_ti(vista):
    @wraps(vista)
    def protegida(request, *args, **kwargs):
        if not PerfilUsuario.objects.filter(user=request.user, rol=PerfilUsuario.ROL_ADMIN_TI).exists():
            return HttpResponseForbidden('Acceso reservado al administrador TI')
        return vista(request, *args, **kwargs)
    return protegida


def fecha(valor):
    if not valor:
        return '-'
    if isinstance(valor, datetime):
        if timezone.is_aware(valor):
            valor = timezone.localtime(valor)
        return valor.strftime('%d/%m/%Y %H:%M')
    if isinstance(valor, date):
        return valor.strftime('%d/%m/%Y')
    return str(valor)


def exportar_reporte(formato, nombre, titulo, secciones):
    if formato == 'excel':
        libro = Workbook()
        for numero, (nombre_hoja, columnas, filas) in enumerate(secciones):
            hoja = libro.active if numero == 0 else libro.create_sheet()
            hoja.title = nombre_hoja[:31]
            hoja.append(columnas)
            for celda in hoja[1]:
                celda.font = Font(bold=True, color='FFFFFF')
                celda.fill = PatternFill('solid', fgColor='253548')
            for fila in filas:
                hoja.append([str(valor) if valor is not None else '' for valor in fila])
                for celda in hoja[hoja.max_row]:
                    celda.data_type = 's'  # Los textos ingresados por usuarios nunca deben convertirse en fórmulas.
            hoja.freeze_panes = 'A2'
            hoja.auto_filter.ref = hoja.dimensions
            for columna in hoja.columns:
                letra = columna[0].column_letter
                hoja.column_dimensions[letra].width = min(42, max(14, *(len(str(c.value or '')) + 2 for c in columna)))
        salida = BytesIO()
        libro.save(salida)
        tipo = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        extension = 'xlsx'
    elif formato == 'pdf':
        salida = BytesIO()
        documento = SimpleDocTemplate(salida, pagesize=landscape(A4), leftMargin=1 * cm,
                                      rightMargin=1 * cm, topMargin=1 * cm, bottomMargin=1 * cm)
        estilos = getSampleStyleSheet()
        celda = estilos['BodyText'].clone('CeldaReporte')
        celda.fontSize = 7
        celda.leading = 9
        encabezado = celda.clone('EncabezadoReporte')
        encabezado.textColor = colors.white
        contenido = [Paragraph(escape(titulo), estilos['Title']), Spacer(1, 0.4 * cm)]
        for nombre_hoja, columnas, filas in secciones:
            contenido.extend([Paragraph(escape(nombre_hoja), estilos['Heading2']), Spacer(1, 0.2 * cm)])
            datos = [[Paragraph(escape(str(valor)), encabezado) for valor in columnas]]
            datos.extend([Paragraph(escape(str(valor if valor is not None else '')), celda)
                          for valor in fila] for fila in filas)
            tabla = LongTable(datos, colWidths=[documento.width / len(columnas)] * len(columnas), repeatRows=1)
            tabla.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#253548')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
                ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                ('GRID', (0, 0), (-1, -1), 0.25, colors.lightgrey),
                ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.whitesmoke]),
            ]))
            contenido.extend([tabla, Spacer(1, 0.5 * cm)])
        documento.build(contenido)
        tipo = 'application/pdf'
        extension = 'pdf'
    else:
        raise ValueError('Formato de reporte no admitido')
    respuesta = HttpResponse(salida.getvalue(), content_type=tipo)
    respuesta['Content-Disposition'] = f'attachment; filename="{nombre}.{extension}"'
    return respuesta
